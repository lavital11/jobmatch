"""Extract plain text from uploaded CV files (PDF, DOCX, TXT), in memory only."""

from __future__ import annotations

import io
import re
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class CVParseError(Exception):
    """Raised when a CV file can't be turned into usable text."""


def normalize_whitespace(text: str) -> str:
    """Collapse all whitespace runs (including line breaks) into single spaces.

    PDF/DOCX extraction frequently breaks lines mid-sentence (wrapping at the
    visual line width rather than at sentence boundaries). Left as-is, that
    makes two semantically identical strings - e.g. the CV text vs. a model's
    quoted "original" excerpt of it - fail a naive exact-match comparison
    purely because of where line breaks happen to fall. Normalizing both
    sides before comparing (or storing) avoids that false mismatch.
    """
    return re.sub(r"\s+", " ", text).strip()


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Return normalized plain text extracted from `file_bytes`, by `filename`'s extension."""
    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        raw = _extract_pdf(file_bytes)
    elif ext == ".docx":
        raw = _extract_docx(file_bytes)
    elif ext == ".txt":
        raw = _extract_txt(file_bytes)
    else:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise CVParseError(
            f"Unsupported file type '{ext or '(no extension)'}'. Supported types: {supported}."
        )

    return normalize_whitespace(raw)


def _extract_pdf(file_bytes: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(file_bytes))
    except Exception as exc:
        raise CVParseError(f"Could not open this PDF file: {exc}") from exc

    pages_text = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(pages_text).strip()
    if not text:
        raise CVParseError(
            "No extractable text found in this PDF. It may be a scanned image rather than "
            "real text - try pasting the CV text instead."
        )
    return text


def _extract_docx(file_bytes: bytes) -> str:
    from docx import Document

    try:
        document = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise CVParseError(f"Could not open this DOCX file: {exc}") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]

    # Many resumes put skills/experience in tables, so pull cell text too.
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    text = "\n".join(parts).strip()
    if not text:
        raise CVParseError("No extractable text found in this DOCX file.")
    return text


def _extract_txt(file_bytes: bytes) -> str:
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise CVParseError(f"Could not decode this file as UTF-8 text: {exc}") from exc

    text = text.strip()
    if not text:
        raise CVParseError("This text file appears to be empty.")
    return text
